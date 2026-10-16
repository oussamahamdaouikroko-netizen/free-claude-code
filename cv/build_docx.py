#!/usr/bin/env python3
"""
Build the editable Word version of the CV, mirroring the PDF design.

- A4, single page, same navy / charcoal / electric-blue palette
- Calibri (the modern default available in every Word install) so the file
  looks right on a recruiter's machine without font installation
- Real clickable hyperlinks (mailto, tel, LinkedIn, portfolio)
- Straightforward paragraph + borderless table structure for clean ATS parsing

Run:  python3 build_docx.py
Out:  Oussama_Hamdaoui_Video_Editor_CV.docx
"""

from __future__ import annotations

import os
import sys

import cv_content as C
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DOCX = os.path.join(HERE, "Oussama_Hamdaoui_Video_Editor_CV.docx")

# ------------------------------------------------------------------ palette
NAVY = RGBColor(0x0B, 0x1B, 0x33)
INK = RGBColor(0x2A, 0x34, 0x41)
MUTED = RGBColor(0x6B, 0x76, 0x88)
ACCENT = RGBColor(0x1F, 0x6F, 0xEB)
RULE = "D8DFEA"

BODY_FONT = "Calibri"
BODY_SIZE = Pt(10.5)

MARGIN = Pt(46)  # left / right  (matches the PDF's 46pt)
MARGIN_TB = Pt(40)  # top / bottom
CONTENT_W = Pt(595.276 - 92)  # A4 width minus both margins


# ------------------------------------------------------------------ helpers
def set_run_font(run, name: str = BODY_FONT) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def style_run(
    run,
    size=BODY_SIZE,
    bold=False,
    italic=False,
    color=INK,
    font=BODY_FONT,
    spacing=None,
):
    set_run_font(run, font)
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    if spacing is not None:  # manual tracking in twentieths of a point
        rpr = run._element.get_or_add_rPr()
        el = OxmlElement("w:spacing")
        el.set(qn("w:val"), str(int(spacing)))
        rpr.append(el)
    return run


def para(
    container, space_before=0, space_after=0, line=None, keep_next=False, align=None
):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if line:
        pf.line_spacing = line
    pf.keep_with_next = keep_next
    pf.widow_control = True
    if align is not None:
        p.alignment = align
    return p


def bottom_border(paragraph, color=RULE, size=6, space=8):
    """Hairline rule under a paragraph (like the PDF section rules)."""
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    ppr.append(borders)


def add_hyperlink(
    paragraph,
    url: str,
    text: str,
    size=BODY_SIZE,
    bold=False,
    color=ACCENT,
    spacing=None,
):
    """Insert a real, clickable hyperlink run into a paragraph."""
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), BODY_FONT)
    rpr.append(rfonts)
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), f"{color}")
    rpr.append(color_el)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size.pt * 2)))
    rpr.append(sz)
    if bold:
        rpr.append(OxmlElement("w:b"))
    if spacing is not None:
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(int(spacing)))
        rpr.append(sp)

    run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def section_heading(doc, text, space_before=15.5, space_after=7.0):
    p = para(doc, space_before=space_before, space_after=space_after, keep_next=True)
    style_run(p.add_run(text.upper()), size=Pt(10), bold=True, color=NAVY, spacing=24)
    bottom_border(p)
    return p


def bullet(doc, text, space_after=3.0):
    p = para(doc, space_after=space_after, line=1.06)
    p.paragraph_format.left_indent = Pt(13)
    p.paragraph_format.first_line_indent = Pt(-13)
    style_run(p.add_run("\u2022  "), color=ACCENT, bold=True)
    # keep **bold** emphasis on tool names
    for chunk, is_bold in split_bold(text):
        style_run(p.add_run(chunk), bold=is_bold, color=INK)
    return p


def split_bold(text: str):
    """Split the inline <b> markup used in cv_content into (chunk, bold)."""
    out, i = [], 0
    while True:
        start = text.find("<b>", i)
        if start == -1:
            out.append((text[i:], False))
            break
        out.append((text[i:start], False))
        end = text.find("</b>", start)
        out.append((text[start + 3 : end], True))
        i = end + 4
    return [(c, b) for c, b in out if c]


def cell_free(table):
    """Remove the default table borders so the grid is invisible."""
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        borders.append(el)
    tbl_pr.append(borders)


# ------------------------------------------------------------------ build
def build(out_path: str = OUT_DOCX) -> str:
    doc = Document()

    # base style
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.font.color.rgb = INK
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), BODY_FONT)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.10

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = MARGIN
    section.right_margin = MARGIN
    section.top_margin = MARGIN_TB
    section.bottom_margin = MARGIN_TB
    section.header_distance = Pt(24)
    section.footer_distance = Pt(24)

    # ---------------------------------------------------------- header
    p = para(doc, space_after=1.5)
    style_run(p.add_run(C.NAME_UPPER), size=Pt(23), bold=True, color=NAVY)

    p = para(doc, space_after=7.0)
    style_run(
        p.add_run(C.JOB_TITLE.upper()), size=Pt(10), bold=True, color=ACCENT, spacing=14
    )

    # contact block: two items per line, tab separated
    def contact_line(items, last=False):
        p = para(doc, space_after=1.0)
        p.paragraph_format.tab_stops.add_tab_stop(Pt(258), WD_TAB_ALIGNMENT.LEFT)
        for idx, (label, text, url) in enumerate(items):
            if idx:
                style_run(p.add_run("\t"), size=Pt(9.5))
            if label:
                style_run(p.add_run(label + " "), size=Pt(9.5), bold=True, color=MUTED)
            if url:
                add_hyperlink(
                    p,
                    url,
                    text,
                    size=Pt(9.5),
                    color=ACCENT
                    if url.startswith(("https", "tel:"))
                    and not url.startswith("mailto")
                    else INK,
                )
                if url.startswith("mailto"):
                    # e-mail reads better in ink than in blue
                    for hyperlink in p._p.findall(qn("w:hyperlink"))[-1:]:
                        for col in hyperlink.iter(qn("w:color")):
                            col.set(qn("w:val"), "2A3441")
            else:
                style_run(p.add_run(text), size=Pt(9.5), color=INK)
        if last:
            bottom_border(p, color="1F6FEB", size=12, space=6)
        return p

    contact_line(
        [
            ("", C.LOCATION, None),
            ("Tel", C.PHONE, C.PHONE_URL),
        ]
    )
    contact_line(
        [
            ("Email", C.EMAIL, f"mailto:{C.EMAIL}"),
            ("LinkedIn", C.LINKEDIN_LABEL, C.LINKEDIN_URL),
        ]
    )
    contact_line([("Portfolio", C.PORTFOLIO_LABEL, C.PORTFOLIO_URL)], last=True)

    # ---------------------------------------------------------- summary
    section_heading(doc, "Professional Summary", space_before=14.0)
    p = para(doc, space_after=0, line=1.12)
    for chunk, is_bold in split_bold(C.SUMMARY_HTML):
        style_run(p.add_run(chunk), bold=is_bold)

    # ---------------------------------------------------------- skills
    section_heading(doc, "Core Skills")
    rows = (len(C.SKILLS) + 1) // 2
    table = doc.add_table(rows=rows, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell_free(table)
    table.autofit = False
    for r in range(rows):
        for cidx in range(2):
            index = cidx * rows + r
            cell = table.cell(r, cidx)
            cell.width = Pt(251)
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_after = Pt(2.8)
            cp.paragraph_format.line_spacing = 1.06
            if index < len(C.SKILLS):
                style_run(cp.add_run("\u25aa  "), color=ACCENT, bold=True)
                style_run(cp.add_run(C.SKILLS[index]))
    # a trailing spacer paragraph keeps the section rhythm
    para(doc, space_after=0)

    # ---------------------------------------------------------- experience
    section_heading(doc, "Professional Experience")
    p = para(doc, space_after=5.0, keep_next=True)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(503), WD_TAB_ALIGNMENT.RIGHT)
    style_run(p.add_run(C.ROLE_TITLE), size=Pt(12), bold=True, color=NAVY)
    style_run(p.add_run("   " + C.ROLE_DESC), size=Pt(11), color=MUTED)
    style_run(p.add_run("\t" + C.ROLE_META), size=Pt(9.5), color=MUTED)

    for b in C.BULLETS_HTML:
        bullet(doc, b)

    # ---------------------------------------------------------- education
    section_heading(doc, "Education")
    p = para(doc, space_after=1.0)
    style_run(p.add_run(C.EDUCATION_DEGREE), size=Pt(11.5), bold=True, color=NAVY)
    style_run(p.add_run("   " + C.EDUCATION_NATIVE), size=Pt(10.2), color=MUTED)
    p = para(doc, space_after=0)
    style_run(p.add_run(C.EDUCATION_NOTE), size=Pt(10), color=MUTED)

    # ---------------------------------------------------------- extras
    section_heading(doc, "Technical Skills")
    p = para(doc, space_after=3.0)
    style_run(p.add_run(C.TECH_SKILLS))

    section_heading(doc, "Languages")
    p = para(doc, space_after=0)
    style_run(p.add_run(C.LANGUAGES_TEXT))

    # ---------------------------------------------------------- metadata
    props = doc.core_properties
    props.title = C.PDF_TITLE
    props.author = C.NAME
    props.subject = C.PDF_SUBJECT
    props.keywords = C.KEYWORDS
    props.comments = (
        "Editable Word version of the one-page CV. Identical content to the PDF export."
    )

    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    path = build()
    print(f"wrote {path} ({os.path.getsize(path):,} bytes)")
    sys.exit(0)
